"""Modality-neutral ingestion, chunking, and hybrid retrieval for project knowledge.

The index stores extracted text and source coordinates, not original files. Original
files remain artifacts in the workspace. Embeddings are optional and supplied by an
adapter; lexical retrieval always works without an external model or dependency.
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import re
import sqlite3
import uuid
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Iterable


@dataclass(frozen=True)
class ExtractedDocument:
    source_id: str
    name: str
    media_type: str
    text: str
    sha256: str
    metadata: dict


@dataclass(frozen=True)
class KnowledgeChunk:
    chunk_id: str
    source_id: str
    source_name: str
    ordinal: int
    text: str
    start_char: int
    end_char: int
    metadata: dict


@dataclass(frozen=True)
class SearchHit:
    chunk: KnowledgeChunk
    score: float
    lexical_score: float
    semantic_score: float


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._hidden += 1
        elif tag.lower() in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._hidden:
            self._hidden -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._hidden:
            self.parts.append(data)


def extract_document(
    name: str,
    content: bytes,
    *,
    media_type: str | None = None,
    max_bytes: int = 50 * 1024 * 1024,
    metadata: dict | None = None,
) -> ExtractedDocument:
    """Extract text from common formats; PDF/DOCX use optional libraries when installed."""
    if not isinstance(content, bytes):
        raise TypeError("content must be bytes")
    if len(content) > max_bytes:
        raise ValueError(f"Document exceeds configured limit ({max_bytes} bytes)")
    suffix = Path(name).suffix.casefold()
    mime = (media_type or "").casefold()
    digest = hashlib.sha256(content).hexdigest()
    text: str
    if suffix in {".txt", ".md", ".rst", ".py", ".js", ".ts", ".tsx", ".jsx", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".log", ".xml"} or mime.startswith("text/"):
        text = content.decode("utf-8-sig", errors="replace")
    elif suffix in {".html", ".htm"} or "html" in mime:
        parser = _TextHTMLParser()
        parser.feed(content.decode("utf-8-sig", errors="replace"))
        text = "".join(parser.parts)
    elif suffix == ".json" or "json" in mime:
        raw = json.loads(content.decode("utf-8-sig"))
        text = json.dumps(raw, ensure_ascii=False, indent=2, sort_keys=True)
    elif suffix == ".csv" or "csv" in mime:
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))
        text = "\n".join(" | ".join(cell for cell in row) for row in rows)
    elif suffix == ".pdf" or "pdf" in mime:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF extraction requires optional dependency 'pypdf'") from exc
        reader = PdfReader(io.BytesIO(content))
        text = "\n\n".join(f"[Page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages))
    elif suffix == ".docx" or "wordprocessingml" in mime:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX extraction requires optional dependency 'python-docx'") from exc
        document = Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(f"Unsupported text extraction format: {name!r}; register a domain extractor for this modality")
    text = text.replace("\x00", "").strip()
    return ExtractedDocument(
        source_id=str(uuid.uuid4()), name=Path(name).name or "unnamed",
        media_type=media_type or _guess_media_type(suffix), text=text, sha256=digest,
        metadata=dict(metadata or {}),
    )


def _guess_media_type(suffix: str) -> str:
    return {
        ".txt": "text/plain", ".md": "text/markdown", ".html": "text/html",
        ".htm": "text/html", ".json": "application/json", ".csv": "text/csv",
        ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }.get(suffix, "text/plain")


def chunk_document(
    document: ExtractedDocument, *, chunk_chars: int = 2400, overlap_chars: int = 250,
) -> list[KnowledgeChunk]:
    """Chunk on paragraph/sentence boundaries where possible while retaining source offsets."""
    if chunk_chars < 100 or overlap_chars < 0 or overlap_chars >= chunk_chars:
        raise ValueError("Require chunk_chars >= 100 and 0 <= overlap_chars < chunk_chars")
    text = document.text
    if not text:
        return []
    chunks: list[KnowledgeChunk] = []
    start = 0
    ordinal = 0
    while start < len(text):
        hard_end = min(len(text), start + chunk_chars)
        end = hard_end
        if hard_end < len(text):
            floor = start + int(chunk_chars * 0.65)
            candidates = [text.rfind("\n\n", floor, hard_end), text.rfind("\n", floor, hard_end), text.rfind(". ", floor, hard_end)]
            boundary = max(candidates)
            if boundary > start:
                end = boundary + (2 if text[boundary:boundary + 2] == ". " else (2 if text[boundary:boundary + 2] == "\n\n" else 1))
        body = text[start:end].strip()
        if body:
            left_trim = len(text[start:end]) - len(text[start:end].lstrip())
            actual_start = start + left_trim
            actual_end = actual_start + len(body)
            chunks.append(KnowledgeChunk(
                chunk_id=f"{document.source_id}:{ordinal}", source_id=document.source_id,
                source_name=document.name, ordinal=ordinal, text=body,
                start_char=actual_start, end_char=actual_end, metadata=dict(document.metadata),
            ))
            ordinal += 1
        if end >= len(text):
            break
        start = max(start + 1, end - overlap_chars)
    return chunks


def _tokens(text: str) -> list[str]:
    return re.findall(r"[\w'-]{2,}", text.casefold())


class KnowledgeIndex:
    """Persistent SQLite chunk index with lexical ranking and optional semantic reranking."""

    def __init__(self, database_path: str | Path):
        self.path = Path(database_path).resolve()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS knowledge_sources (
                    source_id TEXT PRIMARY KEY, name TEXT NOT NULL, media_type TEXT NOT NULL,
                    sha256 TEXT NOT NULL, metadata TEXT NOT NULL, text_length INTEGER NOT NULL
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    chunk_id TEXT PRIMARY KEY, source_id TEXT NOT NULL, ordinal INTEGER NOT NULL,
                    text TEXT NOT NULL, start_char INTEGER NOT NULL, end_char INTEGER NOT NULL,
                    metadata TEXT NOT NULL, FOREIGN KEY(source_id) REFERENCES knowledge_sources(source_id)
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_source ON knowledge_chunks(source_id, ordinal);
            """)

    def _connect(self) -> sqlite3.Connection:
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        return db

    def ingest(self, document: ExtractedDocument, *, chunk_chars: int = 2400, overlap_chars: int = 250) -> int:
        chunks = chunk_document(document, chunk_chars=chunk_chars, overlap_chars=overlap_chars)
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            # Same name and digest is idempotent; changed content replaces the indexed revision.
            old = db.execute("SELECT source_id FROM knowledge_sources WHERE name=? AND sha256=?", (document.name, document.sha256)).fetchone()
            if old:
                return db.execute("SELECT COUNT(*) FROM knowledge_chunks WHERE source_id=?", (old["source_id"],)).fetchone()[0]
            old_rows = db.execute("SELECT source_id FROM knowledge_sources WHERE name=?", (document.name,)).fetchall()
            for old_row in old_rows:
                db.execute("DELETE FROM knowledge_chunks WHERE source_id=?", (old_row["source_id"],))
            db.execute("DELETE FROM knowledge_sources WHERE name=?", (document.name,))
            db.execute(
                "INSERT INTO knowledge_sources VALUES (?, ?, ?, ?, ?, ?)",
                (document.source_id, document.name, document.media_type, document.sha256,
                 json.dumps(document.metadata, ensure_ascii=False), len(document.text)),
            )
            db.executemany(
                "INSERT INTO knowledge_chunks VALUES (?, ?, ?, ?, ?, ?, ?)",
                [(c.chunk_id, c.source_id, c.ordinal, c.text, c.start_char, c.end_char,
                  json.dumps(c.metadata, ensure_ascii=False)) for c in chunks],
            )
        return len(chunks)

    def search(
        self, query: str, *, limit: int = 8,
        semantic_score: Callable[[str, str], float] | None = None,
        semantic_weight: float = 0.35,
    ) -> list[SearchHit]:
        if not query.strip() or limit < 1:
            return []
        if not 0 <= semantic_weight <= 1:
            raise ValueError("semantic_weight must be between 0 and 1")
        with self._connect() as db:
            rows = db.execute(
                "SELECT c.*, s.name AS source_name FROM knowledge_chunks c JOIN knowledge_sources s USING(source_id)"
            ).fetchall()
        qtokens = set(_tokens(query))
        if not qtokens:
            return []
        tokenized = [(row, _tokens(row["text"])) for row in rows]
        document_frequency = {
            term: sum(term in set(tokens) for _, tokens in tokenized)
            for term in qtokens
        }
        raw: list[tuple[sqlite3.Row, float, float]] = []
        for row, tokens in tokenized:
            token_set = set(tokens)
            overlap = len(qtokens & token_set)
            semantic = float(semantic_score(query, row["text"])) if semantic_score else 0.0
            # A semantic adapter may retrieve related chunks with no literal term overlap.
            if not overlap and not semantic_score:
                continue
            tf = sum(min(tokens.count(t), 4) for t in qtokens)
            idf = sum(
                math.log(1 + (len(rows) + 1) / (1 + document_frequency[t]))
                for t in qtokens if t in token_set
            )
            lexical = (tf * idf) / (tf + 1.2 * (0.25 + 0.75 * len(tokens) / 400)) if overlap else 0.0
            raw.append((row, lexical, semantic))
        if not raw:
            return []
        max_lex = max((x[1] for x in raw), default=1) or 1
        # Semantic callbacks are expected to return cosine-like scores; normalize defensively.
        sem_min = min((x[2] for x in raw), default=0)
        sem_max = max((x[2] for x in raw), default=0)
        hits = []
        for row, lexical, semantic in raw:
            lex_norm = lexical / max_lex
            sem_norm = (semantic - sem_min) / (sem_max - sem_min) if sem_max > sem_min else (1.0 if semantic_score and semantic > 0 else 0.0)
            weight = semantic_weight if semantic_score else 0.0
            score = (1 - weight) * lex_norm + weight * sem_norm
            chunk = KnowledgeChunk(
                chunk_id=row["chunk_id"], source_id=row["source_id"], source_name=row["source_name"],
                ordinal=row["ordinal"], text=row["text"], start_char=row["start_char"], end_char=row["end_char"],
                metadata=json.loads(row["metadata"]),
            )
            hits.append(SearchHit(chunk, score, lex_norm, sem_norm))
        return sorted(hits, key=lambda hit: hit.score, reverse=True)[:limit]

    def remove_source(self, source_id: str) -> bool:
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            db.execute("DELETE FROM knowledge_chunks WHERE source_id=?", (source_id,))
            cur = db.execute("DELETE FROM knowledge_sources WHERE source_id=?", (source_id,))
            return cur.rowcount > 0

    def list_sources(self, *, limit: int = 100) -> list[dict]:
        with self._connect() as db:
            rows = db.execute("SELECT source_id,name,media_type,sha256,text_length FROM knowledge_sources ORDER BY name LIMIT ?", (max(1, min(limit, 1000)),)).fetchall()
        return [dict(row) for row in rows]
